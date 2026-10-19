import csv
import io
import json
import time
from datetime import timedelta
from unittest import mock

import openpyxl
from django.conf import settings
from django.contrib.auth.hashers import PBKDF2PasswordHasher, check_password, make_password
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, connection
from django.test import Client, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from core import usuarios_escuela as ue
from core.forms import ConvenioForm, EscuelaForm
from core.models import DocumentoEscuela, Escuela, TipoUsuario, Usuario

FORZAR_CAMBIO = 'core.middleware.ForzarCambioPasswordMiddleware'
CON_CAMBIO_FORZADO = override_settings(MIDDLEWARE=[m for m in settings.MIDDLEWARE if m != FORZAR_CAMBIO] + [FORZAR_CAMBIO])
PASSWORD = 'Test1234!'
NUEVA_PASSWORD = 'NuevaSegura#2026'


def alumno(n=1, **extra):
    datos = {'nombre': f'Alumno{n}', 'apellidos': 'Pérez López', 'grado': 1, 'grupo': 'A'}
    datos.update(extra)
    return datos


def docente(n=1, **extra):
    datos = {'nombre': f'Docente{n}', 'apellidos': 'Gómez Ruiz', 'correo': f'docente{n}@gmail.com', 'telefono': ''}
    datos.update(extra)
    return datos


def excel(filas, encabezados=('Nombre(s)', 'Apellidos', 'Grado', 'Grupo'), nombre='grupo-1A.xlsx'):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(list(encabezados))
    for fila in filas:
        hoja.append(list(fila))
    salida = io.BytesIO()
    libro.save(salida)
    return SimpleUploadedFile(nombre, salida.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class BaseUsuariosEscuela(TestCase):
    """Fixtures compartidas por clase; el hash normal se calcula una sola vez (cuesta ~0.7 s)."""

    @classmethod
    def setUpTestData(cls):
        for id_tipo, nombre in ((1, 'Administrador de sistema'), (3, 'Usuario general'),
                                (4, 'Enlace'), (5, 'Docente'), (6, 'Alumno')):
            TipoUsuario.objects.get_or_create(id_tipo=id_tipo, defaults={'nombre_tipo': nombre})
        cls.hash = make_password(PASSWORD)
        cls.escuela = Escuela.objects.create(nombre='Escuela Uno', municipio='Xalapa', estado='Veracruz', nivel='secundaria')
        cls.otra = Escuela.objects.create(nombre='Escuela Dos', municipio='Xalapa', estado='Veracruz', nivel='superior')
        cls.inactiva = Escuela.objects.create(nombre='Escuela Inactiva', municipio='X', estado='Y', nivel='superior', activa=False)

        cls.enlace = cls.usuario('enlace_u', 4, cls.escuela)
        cls.enlace_otra = cls.usuario('enlace_d', 4, cls.otra)
        cls.enlace_sin_escuela = cls.usuario('enlace_s', 4, None)
        cls.enlace_inactiva = cls.usuario('enlace_i', 4, cls.inactiva)
        cls.docente = cls.usuario('docente_u', 5, cls.escuela)
        cls.alumno = cls.usuario('alumno_u', 6, cls.escuela, email=None, grado=1, grupo='A')
        cls.alumno_otra = cls.usuario('alumno_d', 6, cls.otra, email=None, grado=1, grupo='A')
        cls.general = cls.usuario('general_u', 3, None)
        cls.admin = cls.usuario('admin_u', 1, None, is_superuser=True, is_staff=True)

    @classmethod
    def usuario(cls, login, tipo_id, escuela, **extra):
        datos = {'email': f'{login}@gmail.com', **extra}
        return Usuario.objects.create(
            usuario=login, password=cls.hash, nombre=login.split('_')[0].capitalize(), aPaterno='Prueba',
            aMaterno='', tipousuario_id=tipo_id, escuela=escuela, **datos,
        )

    def setUp(self):
        self.client = Client()

    def entrar(self, usuario):
        self.client.force_login(usuario)

    def post_json(self, nombre_ruta, datos, **kwargs):
        return self.client.post(reverse(nombre_ruta, **kwargs), json.dumps(datos), content_type='application/json')

    def subir(self, archivo):
        return self.client.post(reverse('api-escuelas-alumnos-excel'), {'archivo': archivo})

    def nuevos(self):
        return Usuario.objects.filter(nombre__regex=r'^(Alumno|Docente)[0-9]+$')


# ── Servicio: validación ──────────────────────────────────────────────────────

class ValidacionTest(BaseUsuariosEscuela):

    def test_normaliza_docente(self):
        limpio, errores = ue.validar_persona({
            'nombre': '  María   José ', 'apellidos': ' Hernández  de la Cruz ',
            'correo': ' MJ.Hernandez@Escuela.EDU.MX ', 'telefono': '228 165 0525',
        }, ue.ROL_DOCENTE)
        self.assertEqual(errores, {})
        self.assertEqual(
            (limpio['nombre'], limpio['aPaterno'], limpio['aMaterno'], limpio['correo'], limpio['telefono']),
            ('María José', 'Hernández', 'de la Cruz', 'mj.hernandez@escuela.edu.mx', '228-165-0525'))
        self.assertIsNone(limpio['grado'])

    def test_normaliza_alumno_sin_correo(self):
        for grado in (3, 3.0, '3', ' 3° ', '3º'):
            limpio, errores = ue.validar_persona(alumno(grado=grado, grupo=' b ', correo='ignorado@gmail.com'), ue.ROL_ALUMNO, 3)
            self.assertEqual(errores, {}, grado)
            self.assertEqual((limpio['grado'], limpio['grupo'], limpio['correo']), (3, 'B', None))

    def test_obligatorios_con_mensajes_del_mockup(self):
        _, errores = ue.validar_persona({'nombre': ' ', 'apellidos': '', 'grado': '', 'grupo': None}, ue.ROL_ALUMNO)
        self.assertEqual(errores, {'nombre': 'Falta el nombre', 'apellidos': 'Faltan apellidos',
                                   'grado': 'Falta el grado', 'grupo': 'Falta el grupo'})
        _, errores = ue.validar_persona({'nombre': 'X', 'apellidos': 'Y'}, ue.ROL_DOCENTE)
        self.assertEqual(errores, {'correo': 'Falta el correo'})

    def test_grado_segun_nivel_y_grupo(self):
        self.assertEqual(ue.grado_maximo(self.escuela), 3)
        self.assertEqual(ue.grado_maximo(self.otra), 12)
        for grado in (0, 4, 'tres', True, 2.5):
            _, errores = ue.validar_persona(alumno(grado=grado), ue.ROL_ALUMNO, 3)
            self.assertEqual(errores, {'grado': 'El grado debe ser un número del 1 al 3'}, grado)
        for grupo in ('AB', '1', 'Ñ', '-'):
            _, errores = ue.validar_persona(alumno(grupo=grupo), ue.ROL_ALUMNO, 3)
            self.assertIn('grupo', errores, grupo)

    def test_formatos_invalidos_docente(self):
        casos = {
            'correo': [docente(correo='no-es-correo'), docente(correo='docente@hotmail.com'),
                       docente(correo='a' * 250 + '@gmail.com')],
            'telefono': [docente(telefono='123'), docente(telefono='228-165-05255')],
            'nombre': [docente(nombre='n' * 46)],
            'apellidos': [docente(apellidos='p' * 46 + ' Lopez')],
        }
        for campo, filas in casos.items():
            for fila in filas:
                _, errores = ue.validar_persona(fila, ue.ROL_DOCENTE)
                self.assertIn(campo, errores, fila)

    def test_lista_vacia_no_lista_y_demasiadas_filas(self):
        for filas in ([], None, {'a': 1}, [alumno(i) for i in range(ue.MAX_FILAS_LOTE + 1)]):
            limpias, errores = ue.validar_filas(filas, self.escuela, ue.ROL_ALUMNO)
            self.assertEqual(limpias, [])
            self.assertEqual(len(errores), 1)
            self.assertIsNone(errores[0]['fila'])

    def test_fila_que_no_es_objeto(self):
        _, errores = ue.validar_filas(['texto'], self.escuela, ue.ROL_ALUMNO)
        self.assertEqual(errores[0]['fila'], 1)

    def test_respeta_el_numero_de_fila_del_archivo(self):
        _, errores = ue.validar_filas([alumno(1, fila=2), alumno(2, fila=17, grupo='')], self.escuela, ue.ROL_ALUMNO)
        self.assertEqual(errores, [{'fila': 17, 'campo': 'grupo', 'mensaje': 'Falta el grupo'}])

    def test_correo_docente_repetido_y_registrado(self):
        Usuario.objects.filter(pk=self.alumno_otra.pk).update(email='Ocupado.Correo@gmail.com')
        filas = [docente(1), docente(2, correo='DOCENTE1@gmail.com'), docente(3, correo='ocupado.correo@GMAIL.com')]
        limpias, errores = ue.validar_filas(filas, self.escuela, ue.ROL_DOCENTE)
        self.assertEqual([(e['fila'], e['mensaje']) for e in errores],
                         [(2, ue.MSG_CORREO_REPETIDO), (3, ue.MSG_CORREO_REGISTRADO)])
        self.assertNotIn(self.otra.nombre, errores[1]['mensaje'])
        self.assertEqual([f['fila'] for f in limpias], [1])

    def test_alumno_repetido_y_registrado(self):
        filas = [
            alumno(1, nombre='José', apellidos='Núñez Ávila'),
            alumno(2, nombre='jose', apellidos='NUNEZ avila'),          # mismo alumno sin acentos ni mayúsculas
            alumno(3, nombre='José', apellidos='Núñez Ávila', grupo='B'),  # otro grupo: válido
            alumno(4, nombre='Alumno', apellidos='Prueba'),             # ya registrado (alumno_u, 1° A)
        ]
        _, errores = ue.validar_filas(filas, self.escuela, ue.ROL_ALUMNO)
        self.assertEqual([(e['fila'], e['mensaje']) for e in errores],
                         [(2, ue.MSG_ALUMNO_REPETIDO), (4, ue.MSG_ALUMNO_REGISTRADO)])
        # El mismo nombre en otra escuela no cuenta
        _, errores = ue.validar_filas([alumno(1, nombre='Alumno', apellidos='Prueba')], self.otra, ue.ROL_ALUMNO)
        self.assertEqual([e['mensaje'] for e in errores], [ue.MSG_ALUMNO_REGISTRADO])
        Usuario.objects.filter(pk=self.alumno_otra.pk).update(grupo='C')
        self.assertEqual(ue.validar_filas([alumno(1, nombre='Alumno', apellidos='Prueba')], self.otra, ue.ROL_ALUMNO)[1], [])

    def test_una_sola_consulta(self):
        for rol, fabrica in ((ue.ROL_ALUMNO, alumno), (ue.ROL_DOCENTE, docente)):
            with self.assertNumQueries(1):
                ue.validar_filas([fabrica(i) for i in range(50)], self.escuela, rol)


# ── Servicio: hash, creación y Excel ──────────────────────────────────────────

class HashTemporalTest(BaseUsuariosEscuela):

    def test_hash_ligero_valido_y_rapido(self):
        inicio = time.perf_counter()
        hashes = [ue.hash_password_temporal(f'Temporal{i}') for i in range(50)]
        self.assertLess(time.perf_counter() - inicio, 3)
        self.assertTrue(hashes[0].startswith('pbkdf2_sha256$10000$'))
        self.assertTrue(check_password('Temporal0', hashes[0]))

    def test_login_rehashea_con_iteraciones_normales(self):
        Usuario.objects.filter(pk=self.alumno.pk).update(password=ue.hash_password_temporal('Temporal123'))
        self.assertTrue(self.client.login(username='alumno_u', password='Temporal123'))
        self.alumno.refresh_from_db()
        self.assertEqual(self.alumno.password.split('$')[1], str(PBKDF2PasswordHasher.iterations))


class CrearUsuariosTest(BaseUsuariosEscuela):

    def test_alumnos_sin_correo_con_grado_y_grupo(self):
        limpias, errores = ue.validar_filas([alumno(1), alumno(2, grado=3, grupo='C')], self.escuela, ue.ROL_ALUMNO)
        self.assertEqual(errores, [])
        creados = ue.crear_usuarios(self.escuela, ue.ROL_ALUMNO, limpias)
        # Varios alumnos sin correo conviven con el índice único de email
        alumnos = Usuario.objects.filter(usuario__in=[c['usuario'] for c in creados])
        self.assertEqual(alumnos.filter(email__isnull=True).count(), 2)
        self.assertEqual(sorted(alumnos.values_list('grado', 'grupo')), [(1, 'A'), (3, 'C')])
        self.assertTrue(all(a.fecha_alta for a in alumnos))
        self.assertIsNone(creados[0]['correo'])

    def test_logins_unicos_con_mismo_nombre(self):
        Usuario.objects.create(usuario=f'ana_{self.escuela.id}', email='ana.previa@gmail.com', password='!',
                               nombre='Ana', aPaterno='X', aMaterno='')
        filas = [alumno(i, nombre='Ana Sofía', apellidos=f'Apellido{i}') for i in range(50)]
        limpias, _ = ue.validar_filas(filas, self.escuela, ue.ROL_ALUMNO)
        logins = [c['usuario'] for c in ue.crear_usuarios(self.escuela, ue.ROL_ALUMNO, limpias)]
        self.assertEqual(len(set(logins)), 50)
        self.assertEqual(logins[0], f'ana_{self.escuela.id}_2')

    def test_integrity_error_revierte_todo(self):
        limpias, _ = ue.validar_filas([alumno(i) for i in range(5)], self.escuela, ue.ROL_ALUMNO)
        with mock.patch.object(Usuario.objects, 'bulk_create', side_effect=IntegrityError):
            with self.assertRaises(ue.ConflictoRegistro):
                ue.crear_usuarios(self.escuela, ue.ROL_ALUMNO, limpias)
        self.assertFalse(self.nuevos().exists())

    def test_restablecer_password(self):
        password = ue.restablecer_password_temporal(self.alumno)
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.debe_cambiar_password)
        self.assertTrue(self.alumno.check_password(password))

    def test_plantilla_solo_encabezados(self):
        hoja = openpyxl.load_workbook(io.BytesIO(ue.plantilla_excel())).active
        self.assertEqual([c.value for c in hoja[1]], ['Nombre(s)', 'Apellidos', 'Grado', 'Grupo'])
        self.assertEqual(hoja.max_row, 1)

    def test_contexto_tutorial_sale_de_constantes(self):
        ctx = ue.contexto_tutorial()
        self.assertEqual((ctx['max_filas_lote'], ctx['excel_tamano_maximo_mb']), (ue.MAX_FILAS_LOTE, 5))
        self.assertEqual(ctx['columnas_excel'], ['Nombre(s)', 'Apellidos', 'Grado', 'Grupo'])


class ExcelTest(BaseUsuariosEscuela):

    def setUp(self):
        super().setUp()
        self.entrar(self.enlace)

    def test_vista_previa_como_el_mockup(self):
        r = self.subir(excel([
            ('Ana Sofía', 'López Ruiz', 1, 'A'),
            ('Luis Ángel', 'Pérez Gómez', '1°', 'a'),
            (None, None, None, None),                 # fila vacía: se omite y no corre la numeración
            ('Daniela', 'Ortiz Vega', 1, None),
            ('Mateo', '', 1, 'A'),
        ]))
        self.assertEqual(r.status_code, 200)
        self.assertIn('no-cache', r['Cache-Control'])
        datos = r.json()
        self.assertEqual((datos['archivo'], datos['listas'], datos['con_errores']), ('grupo-1A.xlsx', 2, 2))
        self.assertEqual([(f['fila'], f['errores']) for f in datos['filas']],
                         [(2, []), (3, []), (5, ['Falta el grupo']), (6, ['Faltan apellidos'])])
        self.assertEqual(datos['filas'][1], {'fila': 3, 'nombre': 'Luis Ángel', 'apellidos': 'Pérez Gómez',
                                            'grado': 1, 'grupo': 'A', 'errores': []})
        self.assertFalse(Usuario.objects.filter(nombre__in=['Ana Sofía', 'Luis Ángel']).exists())

    def test_importa_solo_las_filas_listas(self):
        vista = self.subir(excel([('Ana Sofía', 'López Ruiz', 1, 'A'), ('Daniela', 'Ortiz Vega', 1, None),
                                  ('Luis', 'Pérez', 2, 'B')])).json()
        listas = [{k: f[k] for k in ('fila', 'nombre', 'apellidos', 'grado', 'grupo')} for f in vista['filas'] if not f['errores']]
        r = self.post_json('api-escuelas-alumnos-lote', {'filas': listas})
        self.assertEqual(r.status_code, 201)
        self.assertEqual([u['fila'] for u in r.json()['usuarios']], [2, 4])
        self.assertTrue(Usuario.objects.filter(nombre='Ana Sofía', aPaterno='López', aMaterno='Ruiz', grado=1, grupo='A',
                                               escuela=self.escuela, tipousuario_id=6).exists())
        self.assertFalse(Usuario.objects.filter(nombre='Daniela').exists())

    def test_encabezados_sin_importar_mayusculas_ni_acentos(self):
        r = self.subir(excel([('Ana', 'López', 1, 'A')], encabezados=('NOMBRE(S)', 'apellidos', 'Gradó', ' grupo ')))
        self.assertEqual(r.status_code, 200)

    def test_plantilla_descargada_y_llena_se_reconoce(self):
        r = self.client.get(reverse('escuelas-plantilla-alumnos'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheetml', r['Content-Type'])
        self.assertIn('plantilla-alumnos.xlsx', r['Content-Disposition'])
        libro = openpyxl.load_workbook(io.BytesIO(r.content))
        libro.active.append(['José', 'Núñez Ávila', 2, 'B'])
        salida = io.BytesIO()
        libro.save(salida)
        r = self.subir(SimpleUploadedFile('llena.xlsx', salida.getvalue()))
        self.assertEqual(r.json()['listas'], 1)

    def test_archivos_rechazados(self):
        casos = [
            (excel([('Ana', 'López', 1, 'A')], encabezados=('Nombre(s)', 'Apellidos', 'Grado')), 'Falta la columna "Grupo"'),
            (excel([('Ana', 'López', 1, 'A', 'XXXX000000')], encabezados=('Nombre(s)', 'Apellidos', 'Grado', 'Grupo', 'CURP')),
             'La columna "CURP" no es de la plantilla'),
            (excel([]), 'no tiene alumnos'),
            (excel([('Ana', 'López', 1, 'A')], nombre='grupo.csv'), 'formato .xlsx'),
            (SimpleUploadedFile('falso.xlsx', b'no es un zip'), 'no es un Excel'),
            (SimpleUploadedFile('enorme.xlsx', b'0' * (ue.EXCEL_TAMANO_MAXIMO + 1)), 'no debe superar 5 MB'),
            (excel([(f'Alumno{i}', 'López', 1, 'A') for i in range(ue.MAX_FILAS_LOTE + 1)]), 'más de 200 alumnos'),
        ]
        for archivo, mensaje in casos:
            r = self.subir(archivo)
            self.assertEqual(r.status_code, 400, mensaje)
            self.assertIn(mensaje, r.json()['errores'][0]['mensaje'])
        self.assertEqual(self.client.post(reverse('api-escuelas-alumnos-excel')).status_code, 400)

    def test_zip_que_descomprime_demasiado(self):
        with mock.patch.object(ue, 'EXCEL_DESCOMPRIMIDO_MAXIMO', 10):
            r = self.subir(excel([('Ana', 'López', 1, 'A')]))
        self.assertEqual(r.status_code, 400)
        self.assertIn('demasiado grande', r.json()['errores'][0]['mensaje'])


# ── Permisos ──────────────────────────────────────────────────────────────────

PAGINAS_ENLACE = ('escuelas-gestionar-usuarios', 'escuelas-agregar-alumno', 'escuelas-agregar-docente',
                  'escuelas-cargar-alumnos', 'escuelas-tutorial-plantilla', 'escuelas-plantilla-alumnos')


class PermisosTest(BaseUsuariosEscuela):

    def sin_permiso(self):
        return [self.docente, self.alumno, self.general, self.admin, self.enlace_sin_escuela, self.enlace_inactiva]

    def test_endpoints_rechazan_a_quien_no_es_enlace_con_escuela_activa(self):
        for usuario in self.sin_permiso() + [None]:
            self.client.logout()
            if usuario:
                self.entrar(usuario)
            respuestas = [
                self.post_json('api-escuelas-docente-crear', docente(1)),
                self.post_json('api-escuelas-alumno-crear', alumno(1)),
                self.post_json('api-escuelas-alumnos-lote', {'filas': [alumno(1)]}),
                self.subir(excel([('Ana', 'López', 1, 'A')])),
                self.post_json('api-escuelas-usuario-restablecer-password', {}, args=[self.alumno.pk]),
            ]
            for r in respuestas:
                self.assertEqual(r.status_code, 403, usuario)
                self.assertIn('errores', r.json())
        self.assertFalse(self.nuevos().exists())

    def test_endpoints_solo_post(self):
        self.entrar(self.enlace)
        for ruta in ('api-escuelas-docente-crear', 'api-escuelas-alumno-crear', 'api-escuelas-alumnos-excel',
                     'api-escuelas-alumnos-lote'):
            r = self.client.get(reverse(ruta))
            self.assertEqual(r.status_code, 405)
            self.assertIn('no-cache', r['Cache-Control'])

    def test_paginas_solo_para_el_enlace(self):
        for usuario in self.sin_permiso():
            self.entrar(usuario)
            for ruta in PAGINAS_ENLACE:
                self.assertEqual(self.client.get(reverse(ruta)).status_code, 403, (usuario, ruta))
        self.client.logout()
        for ruta in PAGINAS_ENLACE:
            r = self.client.get(reverse(ruta))
            self.assertEqual(r.status_code, 302)
            self.assertIn('/login', r['Location'])
        self.entrar(self.enlace)
        for ruta in PAGINAS_ENLACE:
            self.assertEqual(self.client.get(reverse(ruta)).status_code, 200, ruta)

    def test_docente_y_alumno_solo_consulta(self):
        rutas = [reverse(r) for r in PAGINAS_ENLACE] + [
            reverse('escuelas-crear-enlace'), reverse('escuelas-registrar'),
            reverse('escuelas-editar', args=[self.escuela.id]), reverse('reportes') + '?export=csv',
            reverse('escuelas-gestionar-usuarios') + '?export=csv',
        ]
        for usuario in (self.docente, self.alumno):
            self.entrar(usuario)
            for url in rutas:
                self.assertEqual(self.client.get(url).status_code, 403, (usuario, url))
            r = self.client.post(reverse('escuelas-gestionar-usuarios'), {'action': 'desactivar', 'user_id': self.alumno.pk})
            self.assertEqual(r.status_code, 403)
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.is_active)

    def test_escuela_id_ajeno(self):
        self.entrar(self.enlace)
        self.assertEqual(self.post_json('api-escuelas-docente-crear', docente(1, escuela_id=self.otra.id)).status_code, 403)
        r = self.post_json('api-escuelas-alumnos-lote', {'escuela_id': str(self.otra.id), 'filas': [alumno(2)]})
        self.assertEqual(r.status_code, 403)
        r = self.client.post(reverse('api-escuelas-alumnos-excel'),
                             {'archivo': excel([('Ana', 'López', 1, 'A')]), 'escuela_id': self.otra.id})
        self.assertEqual(r.status_code, 403)
        self.assertFalse(self.nuevos().exists())
        r = self.post_json('api-escuelas-alumnos-lote', {'escuela_id': self.escuela.id, 'filas': [alumno(3)]})
        self.assertEqual(r.status_code, 201)
        self.assertEqual(self.nuevos().get().escuela, self.escuela)

    def test_json_invalido(self):
        self.entrar(self.enlace)
        r = self.client.post(reverse('api-escuelas-alumnos-lote'), 'no-es-json', content_type='application/json')
        self.assertEqual(r.status_code, 400)

    def test_navegacion_mi_escuela_solo_para_el_enlace(self):
        self.entrar(self.enlace)
        r = self.client.get(reverse('escuelas-gestionar-usuarios'))
        self.assertContains(r, 'Mi escuela')
        self.entrar(self.general)
        r = self.client.get(reverse('sobre-nosotros'))
        self.assertNotContains(r, 'Mi escuela')
        self.assertContains(r, reverse('escuelas'))


# ── Registro individual ───────────────────────────────────────────────────────

@CON_CAMBIO_FORZADO
class RegistroIndividualTest(BaseUsuariosEscuela):

    def cambiar_password(self, usuario, temporal):
        """Comprueba el cambio obligatorio y deja al usuario con su contraseña definitiva."""
        r = self.client.get(reverse('escuelas'))
        self.assertRedirects(r, reverse('configuracion', args=[usuario.id]), fetch_redirect_response=False)
        self.client.post(reverse('cambiar-password'), {
            'old_password': temporal, 'new_password1': NUEVA_PASSWORD, 'new_password2': NUEVA_PASSWORD,
        })
        with mock.patch('core.views.get_db_collection') as mongo:
            mongo.return_value = (mock.MagicMock(), mock.MagicMock(find_one=mock.MagicMock(return_value=None)))
            self.assertEqual(self.client.get(reverse('escuelas')).status_code, 200)

    def test_enlace_registra_docente(self):
        self.entrar(self.enlace)
        r = self.post_json('api-escuelas-docente-crear', {
            'nombre': 'Laura', 'apellidos': 'Gómez Ruiz', 'correo': 'Laura.Gomez@itx.edu.mx', 'telefono': '228-111-2233',
        })
        self.assertEqual(r.status_code, 201)
        self.assertIn('no-cache', r['Cache-Control'])
        cred = r.json()['usuarios'][0]
        self.assertEqual(cred['usuario'], f'laura_{self.escuela.id}')
        self.assertEqual(len(cred['password']), 12)

        nuevo = Usuario.objects.get(email='laura.gomez@itx.edu.mx')
        self.assertEqual((nuevo.tipousuario_id, nuevo.escuela_id, nuevo.grado), (5, self.escuela.id, None))
        self.assertEqual((nuevo.aPaterno, nuevo.aMaterno, nuevo.telefono), ('Gómez', 'Ruiz', '228-111-2233'))
        self.assertTrue(nuevo.debe_cambiar_password and nuevo.is_active and nuevo.fecha_alta)
        self.assertIn(nuevo, self.client.get(reverse('escuelas-gestionar-usuarios')).context['usuarios'])

        self.client.logout()
        self.assertTrue(self.client.login(username='laura.gomez@itx.edu.mx', password=cred['password']))
        self.cambiar_password(nuevo, cred['password'])
        self.assertEqual(self.client.get(reverse('escuelas-gestionar-usuarios')).status_code, 403)

    def test_enlace_registra_alumno_que_entra_con_su_usuario(self):
        self.entrar(self.enlace)
        r = self.post_json('api-escuelas-alumno-crear', {'nombre': 'Mateo', 'apellidos': 'Ramírez Torres', 'grado': '3', 'grupo': 'B'})
        self.assertEqual(r.status_code, 201)
        cred = r.json()['usuarios'][0]
        self.assertEqual((cred['grado'], cred['grupo'], cred['correo']), (3, 'B', None))
        nuevo = Usuario.objects.get(usuario=cred['usuario'])
        self.assertEqual((nuevo.tipousuario_id, nuevo.escuela_id, nuevo.email), (6, self.escuela.id, None))

        self.client.logout()
        self.assertTrue(self.client.login(username=cred['usuario'], password=cred['password']))
        self.cambiar_password(nuevo, cred['password'])

    def test_rol_lo_asigna_el_servidor(self):
        self.entrar(self.enlace)
        self.post_json('api-escuelas-alumno-crear', alumno(1, tipousuario_id=4, rol=1))
        self.assertEqual(self.nuevos().get().tipousuario_id, 6)

    def test_errores_por_campo(self):
        self.entrar(self.enlace)
        r = self.post_json('api-escuelas-docente-crear', {'nombre': '', 'apellidos': 'X', 'correo': 'x@hotmail.com'})
        self.assertEqual(r.status_code, 400)
        self.assertEqual({e['campo'] for e in r.json()['errores']}, {'nombre', 'correo'})
        r = self.post_json('api-escuelas-alumno-crear', alumno(1, grado=4))
        self.assertEqual(r.json()['errores'], [{'fila': 1, 'campo': 'grado', 'mensaje': 'El grado debe ser un número del 1 al 3'}])
        r = self.post_json('api-escuelas-alumno-crear', alumno(1, nombre='Alumno', apellidos='Prueba'))
        self.assertEqual(r.json()['errores'][0]['mensaje'], ue.MSG_ALUMNO_REGISTRADO)
        self.assertFalse(self.nuevos().exists())


# ── Lote (C) ──────────────────────────────────────────────────────────────────

class LoteTest(BaseUsuariosEscuela):

    def setUp(self):
        super().setUp()
        self.entrar(self.enlace)

    def test_50_alumnos_rapido_sin_duplicados_ni_huerfanos(self):
        filas = [alumno(i, apellidos=f'López{i}') for i in range(50)]
        inicio = time.perf_counter()
        with CaptureQueriesContext(connection) as consultas:
            r = self.post_json('api-escuelas-alumnos-lote', {'filas': filas})
        self.assertLess(time.perf_counter() - inicio, 5)
        self.assertLessEqual(len(consultas), 12)
        self.assertEqual(r.status_code, 201)
        datos = r.json()
        self.assertEqual(datos['creados'], 50)

        creados = self.nuevos()
        self.assertEqual(creados.filter(tipousuario_id=6, escuela=self.escuela, debe_cambiar_password=True,
                                        email__isnull=True).count(), 50)
        self.assertEqual(len(set(creados.values_list('usuario', flat=True))), 50)
        self.assertFalse(Usuario.objects.filter(tipousuario_id__in=[5, 6], escuela__isnull=True).exists())

        cred = datos['usuarios'][7]
        self.client.logout()
        self.assertTrue(self.client.login(username=cred['usuario'], password=cred['password']))

    def test_una_fila_invalida_no_crea_ninguna(self):
        filas = [alumno(i, apellidos=f'López{i}') for i in range(50)]
        filas[30]['grupo'] = ''
        r = self.post_json('api-escuelas-alumnos-lote', {'filas': filas})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(r.json()['errores'][0]['fila'], 31)
        self.assertFalse(self.nuevos().exists())

    def test_conflicto_por_carrera_devuelve_409(self):
        with mock.patch.object(Usuario.objects, 'bulk_create', side_effect=IntegrityError):
            r = self.post_json('api-escuelas-alumnos-lote', {'filas': [alumno(1), alumno(2)]})
        self.assertEqual(r.status_code, 409)
        self.assertFalse(self.nuevos().exists())

    def test_limites(self):
        for cuerpo in ({'filas': []}, {'filas': [alumno(i) for i in range(201)]}, {}):
            self.assertEqual(self.post_json('api-escuelas-alumnos-lote', cuerpo).status_code, 400)
        self.assertFalse(self.nuevos().exists())


# ── Mi escuela: lista, filtros, exportación y acciones ────────────────────────

class ListaTest(BaseUsuariosEscuela):

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        # Escuela: 117 + docente_u + alumno_u + maria_x = 120 (60 docentes, 60 alumnos, 12 inactivos)
        nuevos = []
        for i in range(117):
            es_docente = i % 2 == 0
            nuevos.append(Usuario(
                usuario=f'lista{i}', email=f'lista{i}@gmail.com' if es_docente else None, password='!',
                nombre=f'Nombre{i:03d}', aPaterno=f'Paterno{i:03d}', aMaterno='',
                tipousuario_id=5 if es_docente else 6, escuela=cls.escuela, is_active=i % 10 != 0,
                grado=None if es_docente else (i % 3) + 1, grupo=None if es_docente else 'AB'[(i // 2) % 2],
            ))
        nuevos.append(Usuario(usuario='maria_x', email=None, password='!', nombre='María', aPaterno='López',
                              aMaterno='Ruiz', tipousuario_id=6, escuela=cls.escuela, grado=2, grupo='C'))
        # Otra escuela: 29 + alumno_d = 30, con el mismo nombre para comprobar que no se mezclan
        nuevos += [
            Usuario(usuario=f'otra{i}', email=None, password='!', nombre='María', aPaterno='López', aMaterno='',
                    tipousuario_id=6, escuela=cls.otra, grado=2, grupo='C')
            for i in range(29)
        ]
        Usuario.objects.bulk_create(nuevos)

    def setUp(self):
        super().setUp()
        self.entrar(self.enlace)

    def lista(self, **params):
        r = self.client.get(reverse('escuelas-gestionar-usuarios'), params)
        self.assertEqual(r.status_code, 200)
        return r.context

    def test_solo_su_escuela_y_paginacion(self):
        ctx = self.lista()
        self.assertEqual(ctx['total_resultados'], 120)
        self.assertEqual((ctx['total_docentes'], ctx['total_alumnos']), (60, 60))
        self.assertEqual(ctx['pagina'].paginator.num_pages, 5)
        self.assertEqual(len(ctx['usuarios']), 25)
        self.assertEqual(len(self.lista(page=5)['usuarios']), 20)
        self.assertEqual(self.lista(page=99)['pagina'].number, 5)
        self.assertEqual(self.lista(page='abc')['pagina'].number, 1)

        vistos = set()
        for n in range(1, 6):
            vistos |= {u.pk for u in self.lista(page=n)['usuarios']}
        self.assertEqual(len(vistos), 120)
        self.assertFalse(Usuario.objects.filter(pk__in=vistos).exclude(escuela=self.escuela).exists())
        self.assertNotIn(self.enlace.pk, vistos)

    def test_alumnos_primero(self):
        usuarios = list(self.lista()['usuarios'])
        self.assertTrue(all(u.tipousuario_id == 6 for u in usuarios))

    def test_busqueda(self):
        self.assertEqual(self.lista(q='Nombre005')['total_resultados'], 1)
        self.assertEqual(self.lista(q='paterno010')['total_resultados'], 1)
        self.assertEqual(self.lista(q='lista6@gmail')['total_resultados'], 1)
        self.assertEqual([u.usuario for u in self.lista(q='maría lópez')['usuarios']], ['maria_x'])
        self.assertEqual(self.lista(q='María Inexistente')['total_resultados'], 0)

    def test_filtros_y_persistencia(self):
        self.assertEqual(self.lista(rol='docente')['total_resultados'], 60)
        self.assertEqual(self.lista(rol='alumno')['total_resultados'], 60)
        self.assertEqual(self.lista(estado='inactivos')['total_resultados'], 12)
        self.assertEqual(self.lista(rol='otro', estado='raro', grado='x', grupo='zz')['total_resultados'], 120)
        esperados = Usuario.objects.filter(escuela=self.escuela, tipousuario_id=6, grado=2, grupo='C').count()
        self.assertEqual(self.lista(grado=2, grupo='c')['total_resultados'], esperados)

        ctx = self.lista()
        self.assertEqual(ctx['grados_disponibles'], [1, 2, 3])
        self.assertEqual(ctx['grupos_disponibles'], ['A', 'B', 'C'])

        ctx = self.lista(q='Nombre', rol='alumno', grado=1, grupo='A', page=2)
        self.assertEqual(ctx['filtros_querystring'], 'q=Nombre&rol=alumno&grado=1&grupo=A')
        self.assertEqual(self.lista()['filtros_querystring'], '')
        # Los enlaces de paginación conservan los filtros
        r = self.client.get(reverse('escuelas-gestionar-usuarios'), {'rol': 'docente'})
        self.assertContains(r, '?rol=docente&amp;page=2')
        self.assertContains(r, '?rol=docente&amp;export=csv')

    def test_exportar_csv_filtrado_y_seguro(self):
        Usuario.objects.filter(usuario='lista1').update(nombre='=HYPERLINK("x")')
        r = self.client.get(reverse('escuelas-gestionar-usuarios'), {'rol': 'alumno', 'export': 'csv'})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r['Content-Type'].startswith('text/csv'))
        contenido = r.content.decode('utf-8')
        self.assertTrue(contenido.startswith('﻿'))
        filas = list(csv.reader(io.StringIO(contenido.lstrip('﻿'))))
        self.assertEqual(filas[0][:4], ['Nombre(s)', 'Apellido paterno', 'Apellido materno', 'Tipo'])
        self.assertEqual(len(filas) - 1, 60)
        self.assertTrue(all(f[3] == 'Alumno' for f in filas[1:]))
        self.assertIn("'=HYPERLINK", contenido)
        self.assertNotIn('otra0', contenido)

    def test_datos_de_la_escuela(self):
        Escuela.objects.filter(pk=self.escuela.pk).update(aviso_privacidad_aceptado='2026-10-04')
        DocumentoEscuela.objects.create(escuela=self.escuela, archivo='x.pdf', fecha_firma='2026-01-10',
                                        fecha_vigencia='2027-12-31')
        r = self.client.get(reverse('escuelas-gestionar-usuarios'))
        self.assertContains(r, 'vigente hasta 31/12/2027')
        self.assertContains(r, 'aceptado 04/10/2026')
        self.assertEqual(r.context['total_modulos'], 0)
        DocumentoEscuela.objects.filter(escuela=self.escuela).update(fecha_vigencia=timezone.localdate() - timedelta(days=1))
        self.assertContains(self.client.get(reverse('escuelas-gestionar-usuarios')), 'vencido el')

    def test_fecha_de_alta(self):
        r = self.client.get(reverse('escuelas-gestionar-usuarios'), {'q': 'maria_x'})
        self.assertContains(r, timezone.localdate().strftime('%d/%m/%Y'))
        Usuario.objects.filter(usuario='maria_x').update(fecha_alta=None)
        r = self.client.get(reverse('escuelas-gestionar-usuarios'), {'q': 'maria_x'})
        self.assertNotContains(r, timezone.localdate().strftime('%d/%m/%Y'))


class AccionesTest(BaseUsuariosEscuela):

    def setUp(self):
        super().setUp()
        self.entrar(self.enlace)

    def accion(self, action, usuario, url=None):
        return self.client.post(url or reverse('escuelas-gestionar-usuarios'), {'action': action, 'user_id': usuario.pk})

    def test_desactivar_y_activar_propios_conserva_filtros(self):
        url = reverse('escuelas-gestionar-usuarios') + '?rol=alumno&page=1'
        r = self.accion('desactivar', self.alumno, url)
        self.assertRedirects(r, url, fetch_redirect_response=False)
        self.alumno.refresh_from_db()
        self.assertFalse(self.alumno.is_active)
        self.accion('activar', self.alumno)
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.is_active)

    def test_no_puede_desactivarse_ni_tocar_otra_escuela(self):
        for usuario in (self.enlace, self.alumno_otra, self.enlace_otra):
            self.accion('desactivar', usuario)
            usuario.refresh_from_db()
            self.assertTrue(usuario.is_active, usuario)
        self.client.post(reverse('escuelas-gestionar-usuarios'), {'action': 'desactivar', 'user_id': 'abc'})

    def test_alta_manual_anterior_ya_no_existe(self):
        self.client.post(reverse('escuelas-gestionar-usuarios'), {
            'action': 'crear', 'usuario': 'intruso', 'email': 'intruso@gmail.com',
            'nombre': 'X', 'aPaterno': 'Y', 'password': 'Pass1234!', 'tipo_usuario': 1,
        })
        self.assertFalse(Usuario.objects.filter(usuario='intruso').exists())

    def test_restablecer_password_solo_propios(self):
        r = self.post_json('api-escuelas-usuario-restablecer-password', {}, args=[self.alumno.pk])
        self.assertEqual(r.status_code, 200)
        self.assertIn('no-cache', r['Cache-Control'])
        datos = r.json()['usuario']
        self.assertEqual((datos['usuario'], datos['correo']), ('alumno_u', None))
        self.alumno.refresh_from_db()
        self.assertTrue(self.alumno.debe_cambiar_password)
        self.assertTrue(self.alumno.check_password(datos['password']))
        for usuario in (self.enlace, self.alumno_otra, self.enlace_otra):
            r = self.post_json('api-escuelas-usuario-restablecer-password', {}, args=[usuario.pk])
            self.assertEqual(r.status_code, 404, usuario)


# ── Pantallas ─────────────────────────────────────────────────────────────────

class PantallasTest(BaseUsuariosEscuela):

    def setUp(self):
        super().setUp()
        self.entrar(self.enlace)

    def test_agregar_alumno_grados_segun_nivel(self):
        r = self.client.get(reverse('escuelas-agregar-alumno'))
        self.assertContains(r, 'Solo se registran nombre, grado y grupo.')
        self.assertEqual(list(r.context['grados']), [1, 2, 3])
        self.assertContains(r, reverse('api-escuelas-alumno-crear'))
        self.assertNotContains(r, 'name="correo"')

    def test_agregar_docente_pide_correo(self):
        r = self.client.get(reverse('escuelas-agregar-docente'))
        self.assertContains(r, 'name="correo"')
        self.assertContains(r, reverse('api-escuelas-docente-crear'))
        self.assertNotContains(r, 'name="grado"')

    def test_carga_y_tutorial_con_limites_del_servidor(self):
        r = self.client.get(reverse('escuelas-cargar-alumnos'))
        self.assertContains(r, f'data-max-bytes="{ue.EXCEL_TAMANO_MAXIMO}"')
        self.assertContains(r, 'hasta 5 MB')
        self.assertContains(r, reverse('escuelas-tutorial-plantilla'))
        r = self.client.get(reverse('escuelas-tutorial-plantilla'))
        self.assertContains(r, f'Hasta {ue.MAX_FILAS_LOTE} alumnos')
        self.assertContains(r, 'hasta 3)')
        self.assertContains(r, reverse('escuelas-plantilla-alumnos'))


# ── Fechas nuevas de la escuela ───────────────────────────────────────────────

class FechasEscuelaFormTest(BaseUsuariosEscuela):

    def test_vigencia_posterior_a_la_firma(self):
        form = ConvenioForm(data={'tipo': 'convenio', 'fecha_firma': '2026-01-10', 'fecha_vigencia': '2026-01-10'},
                            files={'archivo': SimpleUploadedFile('c.pdf', b'%PDF-1.4\n', content_type='application/pdf')})
        self.assertIn('fecha_vigencia', form.errors)
        form = ConvenioForm(data={'tipo': 'convenio', 'fecha_firma': '2026-01-10', 'fecha_vigencia': ''},
                            files={'archivo': SimpleUploadedFile('c.pdf', b'%PDF-1.4\n', content_type='application/pdf')})
        self.assertNotIn('fecha_vigencia', form.errors)

    def test_aviso_de_privacidad_no_futuro(self):
        manana = (timezone.localdate() + timedelta(days=1)).isoformat()
        self.assertIn('aviso_privacidad_aceptado', EscuelaForm(data={'aviso_privacidad_aceptado': manana}).errors)
        self.assertNotIn('aviso_privacidad_aceptado', EscuelaForm(data={'aviso_privacidad_aceptado': ''}).errors)
