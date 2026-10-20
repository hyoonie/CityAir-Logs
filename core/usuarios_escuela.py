"""Alta y gestión de docentes y alumnos por el Enlace de una escuela.

Docentes: nombre, apellidos, correo (su acceso) y teléfono opcional.
Alumnos: nombre, apellidos, grado y grupo; sin correo, entran con el usuario autogenerado.
Las vías de registro (individual y carga desde Excel) usan estas mismas reglas; las vistas solo
resuelven permisos, leen la petición y llaman a este módulo.
"""
import csv
import io
import re
import zipfile

from django.contrib.auth.hashers import PBKDF2PasswordHasher
from django.core.exceptions import ValidationError
from django.core.validators import EmailValidator
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.db.models.functions import Lower
from django.utils import timezone

from .models import Usuario
from .validators import (
    _sin_acentos, dividir_apellidos, generar_password_temporal, generar_usuario_enlace,
    normalizar_telefono, validar_correo_dominio,
)

ROL_ENLACE = 4
ROL_DOCENTE = 5
ROL_ALUMNO = 6
ROLES_GESTIONABLES = (ROL_DOCENTE, ROL_ALUMNO)
CAMPOS_POR_ROL = {
    ROL_DOCENTE: ('nombre', 'apellidos', 'correo', 'telefono'),
    ROL_ALUMNO: ('nombre', 'apellidos', 'grado', 'grupo'),
}

MAX_FILAS_LOTE = 200
USUARIOS_POR_PAGINA = 25

# Grado más alto aceptado según el nivel de la escuela (años o semestres)
GRADO_MAXIMO_POR_NIVEL = {'secundaria': 3, 'media_superior': 6, 'superior': 12}
GRADO_MAXIMO_OTRO = 6
GRUPOS = 'ABCDEFGHIJKLMNOPQRSTUVWXYZ'

# Excel: se lee en el servidor. El límite descomprimido evita archivos "bomba" (un .xlsx es un zip)
EXCEL_TAMANO_MAXIMO = 5 * 1024 * 1024
EXCEL_DESCOMPRIMIDO_MAXIMO = 50 * 1024 * 1024
COLUMNAS_EXCEL = (('nombre', 'Nombre(s)'), ('apellidos', 'Apellidos'), ('grado', 'Grado'), ('grupo', 'Grupo'))
ENCABEZADOS_ACEPTADOS = {
    'nombre': 'nombre', 'nombres': 'nombre', 'apellido': 'apellidos', 'apellidos': 'apellidos',
    'grado': 'grado', 'grupo': 'grupo',
}

FILTROS_ROL = {'docente': ROL_DOCENTE, 'alumno': ROL_ALUMNO}
FILTROS_ESTADO = {'activos': True, 'inactivos': False}

MSG_CORREO_REPETIDO = 'Correo repetido en la lista'
MSG_CORREO_REGISTRADO = 'Ya existe un usuario registrado con este correo'
MSG_ALUMNO_REPETIDO = 'Alumno repetido en el archivo'
MSG_ALUMNO_REGISTRADO = 'Ya existe un alumno con ese nombre en el mismo grado y grupo'

_validar_formato_correo = EmailValidator(message='Correo no válido')


class ConflictoRegistro(Exception):
    """Otro registro ocupó un correo o usuario entre la validación y el guardado."""


class ArchivoInvalido(Exception):
    """El archivo de Excel no se puede procesar; el mensaje se muestra tal cual al Enlace."""


class HashTemporal(PBKDF2PasswordHasher):
    """PBKDF2 ligero solo para contraseñas temporales (aleatorias, 12 caracteres, un solo uso).
    Se usa directo, sin registrarlo en PASSWORD_HASHERS: el hash queda como pbkdf2_sha256 con
    menos iteraciones y Django lo vuelve a hashear con las normales en el primer ingreso."""
    iterations = 10_000


_hash_temporal = HashTemporal()


def hash_password_temporal(password):
    return _hash_temporal.encode(password, _hash_temporal.salt())


# ── Permisos ──────────────────────────────────────────────────────────────────

def enlace_de_escuela(user):
    """Escuela activa del Enlace autenticado, o None si el usuario no puede gestionarla."""
    if not user.is_authenticated or user.tipousuario_id != ROL_ENLACE or not user.escuela_id:
        return None
    escuela = user.escuela
    return escuela if escuela.activa else None


def escuela_ajena(datos, escuela):
    """True si la petición trae un escuela_id distinto del de la escuela del Enlace."""
    if not isinstance(datos, dict):
        return False
    valor = datos.get('escuela_id')
    return valor not in (None, '') and str(valor).strip() != str(escuela.id)


def grado_maximo(escuela):
    return GRADO_MAXIMO_POR_NIVEL.get(escuela.nivel, GRADO_MAXIMO_OTRO)


# ── Validación ────────────────────────────────────────────────────────────────

def _texto(datos, campo):
    valor = datos.get(campo)
    return '' if valor is None else ' '.join(str(valor).split())


def _validar_grado(valor, maximo):
    """Acepta 3, 3.0, '3' o '3°'. Devuelve el entero o lanza ValidationError."""
    fuera_de_rango = ValidationError(f'El grado debe ser un número del 1 al {maximo}')
    if isinstance(valor, bool):
        raise fuera_de_rango
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = re.sub(r'[°º\s]', '', '' if valor is None else str(valor))
    if not texto:
        raise ValidationError('Falta el grado')
    if not texto.isdigit() or not 1 <= int(texto) <= maximo:
        raise fuera_de_rango
    return int(texto)


def _validar_grupo(valor):
    texto = '' if valor is None else str(valor).strip().upper()
    if not texto:
        raise ValidationError('Falta el grupo')
    if len(texto) != 1 or texto not in GRUPOS:
        raise ValidationError('El grupo debe ser una letra (A, B, C…)')
    return texto


def validar_persona(datos, rol, maximo_grado=GRADO_MAXIMO_OTRO):
    """Valida y normaliza una persona sin tocar la BD. Devuelve (limpios, errores_por_campo)."""
    if not isinstance(datos, dict):
        return None, {None: 'La fila no tiene un formato válido'}

    errores = {}
    nombre = _texto(datos, 'nombre')
    apellidos = _texto(datos, 'apellidos')
    paterno, materno = dividir_apellidos(apellidos)

    if not nombre:
        errores['nombre'] = 'Falta el nombre'
    elif len(nombre) > 45:
        errores['nombre'] = 'El nombre debe tener como máximo 45 caracteres'
    if not apellidos:
        errores['apellidos'] = 'Faltan apellidos'
    elif len(paterno) > 45 or len(materno) > 45:
        errores['apellidos'] = 'Cada apellido debe tener como máximo 45 caracteres'

    limpio = {'nombre': nombre, 'aPaterno': paterno, 'aMaterno': materno,
              'correo': None, 'telefono': None, 'grado': None, 'grupo': None}

    if rol == ROL_DOCENTE:
        correo = _texto(datos, 'correo').lower()
        if not correo:
            errores['correo'] = 'Falta el correo'
        elif len(correo) > 255:
            errores['correo'] = 'El correo debe tener como máximo 255 caracteres'
        else:
            try:
                _validar_formato_correo(correo)
                validar_correo_dominio(correo)
            except ValidationError as e:
                errores['correo'] = e.messages[0]
        limpio['correo'] = correo
        telefono = _texto(datos, 'telefono')
        if telefono:
            try:
                limpio['telefono'] = normalizar_telefono(telefono)
            except ValidationError as e:
                errores['telefono'] = e.messages[0]
    else:
        for campo, validar in (('grado', lambda v: _validar_grado(v, maximo_grado)), ('grupo', _validar_grupo)):
            try:
                limpio[campo] = validar(datos.get(campo))
            except ValidationError as e:
                errores[campo] = e.messages[0]

    return (None, errores) if errores else (limpio, {})


def _clave_alumno(nombre, paterno, materno, grado, grupo):
    nombre_completo = ' '.join(f'{nombre} {paterno} {materno}'.split())
    return (_sin_acentos(nombre_completo).casefold(), grado, grupo)


def _numero_fila(datos, posicion):
    """Usa el número de fila del archivo si viene (p. ej. la fila de Excel); si no, la posición."""
    fila = datos.get('fila') if isinstance(datos, dict) else None
    return fila if isinstance(fila, int) and not isinstance(fila, bool) and fila >= 1 else posicion


def validar_filas(filas, escuela, rol, excluir_pk=None):
    """Valida un arreglo de personas con una sola consulta a la BD.
    `excluir_pk` deja fuera de los duplicados al usuario que se está editando.
    Devuelve (filas_limpias, errores); errores es [{fila, campo, mensaje}]."""
    if not isinstance(filas, list):
        return [], [{'fila': None, 'campo': None, 'mensaje': 'Las filas deben enviarse como una lista'}]
    if not filas:
        return [], [{'fila': None, 'campo': None, 'mensaje': 'No hay filas para registrar'}]
    if len(filas) > MAX_FILAS_LOTE:
        return [], [{'fila': None, 'campo': None,
                     'mensaje': f'Se pueden registrar como máximo {MAX_FILAS_LOTE} personas por carga'}]

    maximo = grado_maximo(escuela)
    limpias, errores, vistos = [], [], set()
    for posicion, datos in enumerate(filas, start=1):
        numero = _numero_fila(datos, posicion)
        limpio, errores_fila = validar_persona(datos, rol, maximo)
        errores.extend({'fila': numero, 'campo': c, 'mensaje': m} for c, m in errores_fila.items())
        if limpio is None:
            continue
        if rol == ROL_DOCENTE:
            clave, campo, repetido = limpio['correo'], 'correo', MSG_CORREO_REPETIDO
        else:
            clave = _clave_alumno(limpio['nombre'], limpio['aPaterno'], limpio['aMaterno'], limpio['grado'], limpio['grupo'])
            campo, repetido = 'nombre', MSG_ALUMNO_REPETIDO
        if clave in vistos:
            errores.append({'fila': numero, 'campo': campo, 'mensaje': repetido})
            continue
        vistos.add(clave)
        limpias.append({'fila': numero, 'clave': clave, **limpio})

    if limpias:
        otros = Usuario.objects.exclude(pk=excluir_pk) if excluir_pk else Usuario.objects.all()
        if rol == ROL_DOCENTE:
            registrados = set(
                otros.annotate(correo=Lower('email'))
                .filter(correo__in=vistos).values_list('correo', flat=True)
            )
            campo, mensaje = 'correo', MSG_CORREO_REGISTRADO
        else:
            existentes = otros.filter(
                escuela=escuela, tipousuario_id=ROL_ALUMNO,
                grado__in={f['grado'] for f in limpias}, grupo__in={f['grupo'] for f in limpias},
            ).values_list('nombre', 'aPaterno', 'aMaterno', 'grado', 'grupo')
            registrados = {_clave_alumno(*e) for e in existentes}
            campo, mensaje = 'nombre', MSG_ALUMNO_REGISTRADO
        for limpio in limpias:
            if limpio['clave'] in registrados:
                errores.append({'fila': limpio['fila'], 'campo': campo, 'mensaje': mensaje})

    orden = CAMPOS_POR_ROL[rol]
    errores.sort(key=lambda e: (e['fila'] or 0, orden.index(e['campo']) if e['campo'] in orden else -1))
    con_error = {e['fila'] for e in errores}
    return [f for f in limpias if f['fila'] not in con_error], errores


# ── Creación ──────────────────────────────────────────────────────────────────

def crear_usuarios(escuela, rol_id, filas_limpias):
    """Crea todas las personas en una transacción (todo o nada) con un solo INSERT.
    Devuelve las credenciales para mostrarlas una vez; las contraseñas no se guardan en claro."""
    # Los logins de la escuela terminan en _<id> o _<id>_<n>; se precargan para no consultar uno por uno
    ocupados = {
        u.lower() for u in
        Usuario.objects.filter(usuario__icontains=f'_{escuela.id}').values_list('usuario', flat=True)
    }
    nuevos, credenciales = [], []
    for fila in filas_limpias:
        usuario = generar_usuario_enlace(fila['nombre'], escuela.id, lambda u: u.lower() in ocupados)
        ocupados.add(usuario.lower())
        password = generar_password_temporal()
        nuevos.append(Usuario(
            usuario=usuario, email=fila['correo'], password=hash_password_temporal(password),
            nombre=fila['nombre'], aPaterno=fila['aPaterno'], aMaterno=fila['aMaterno'],
            telefono=fila['telefono'], grado=fila['grado'], grupo=fila['grupo'],
            tipousuario_id=rol_id, escuela=escuela, is_active=True, debe_cambiar_password=True,
        ))
        credenciales.append({
            'fila': fila['fila'],
            'nombre': ' '.join(p for p in (fila['nombre'], fila['aPaterno'], fila['aMaterno']) if p),
            'correo': fila['correo'], 'grado': fila['grado'], 'grupo': fila['grupo'],
            'usuario': usuario, 'password': password,
        })
    try:
        with transaction.atomic():
            Usuario.objects.bulk_create(nuevos)
    except IntegrityError as e:
        raise ConflictoRegistro from e
    return credenciales


CAMPOS_EDITABLES = {
    ROL_DOCENTE: ('nombre', 'aPaterno', 'aMaterno', 'email', 'telefono'),
    ROL_ALUMNO: ('nombre', 'aPaterno', 'aMaterno', 'grado', 'grupo'),
}


def datos_edicion(usuario):
    """Valores actuales en el formato del formulario (apellidos juntos, como en el alta)."""
    return {
        'nombre': usuario.nombre,
        'apellidos': ' '.join(p for p in (usuario.aPaterno, usuario.aMaterno) if p),
        'correo': usuario.email or '', 'telefono': usuario.telefono or '',
        'grado': usuario.grado or '', 'grupo': usuario.grupo or '',
    }


def actualizar_usuario(usuario, escuela, datos):
    """Corrige los datos del registro con las mismas reglas del alta. Solo toca los campos del
    rol: rol, escuela, usuario, contraseña y estado no cambian. Devuelve la lista de errores."""
    rol = usuario.tipousuario_id
    entrada = {campo: datos.get(campo) for campo in CAMPOS_POR_ROL[rol]} if isinstance(datos, dict) else datos
    limpias, errores = validar_filas([entrada], escuela, rol, excluir_pk=usuario.pk)
    if errores:
        return errores
    limpio = limpias[0]
    usuario.nombre, usuario.aPaterno, usuario.aMaterno = limpio['nombre'], limpio['aPaterno'], limpio['aMaterno']
    if rol == ROL_DOCENTE:
        usuario.email, usuario.telefono = limpio['correo'], limpio['telefono']
    else:
        usuario.grado, usuario.grupo = limpio['grado'], limpio['grupo']
    try:
        with transaction.atomic():
            usuario.save(update_fields=CAMPOS_EDITABLES[rol])
    except IntegrityError as e:
        raise ConflictoRegistro from e
    return []


def restablecer_password_temporal(usuario):
    """Asigna una nueva contraseña temporal y la devuelve para mostrarla una sola vez."""
    password = generar_password_temporal()
    usuario.password = hash_password_temporal(password)
    usuario.debe_cambiar_password = True
    usuario.save(update_fields=['password', 'debe_cambiar_password'])
    return password


# ── Excel ─────────────────────────────────────────────────────────────────────

def _celda(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return ' '.join(str(valor).split())


def _encabezado(valor):
    return re.sub(r'[^a-z]', '', _sin_acentos(_celda(valor)).lower())


def leer_excel_alumnos(archivo):
    """Lee la primera hoja del .xlsx. Devuelve [{fila, nombre, apellidos, grado, grupo}] con el
    número de fila de Excel, sin filas vacías. Lanza ArchivoInvalido con un mensaje para el Enlace."""
    import openpyxl  # solo se carga al procesar un archivo

    if not (archivo.name or '').lower().endswith('.xlsx'):
        raise ArchivoInvalido('El archivo debe tener formato .xlsx.')
    if archivo.size > EXCEL_TAMANO_MAXIMO:
        raise ArchivoInvalido(f'El archivo no debe superar {EXCEL_TAMANO_MAXIMO // (1024 * 1024)} MB.')
    try:
        with zipfile.ZipFile(archivo) as contenido:
            if sum(i.file_size for i in contenido.infolist()) > EXCEL_DESCOMPRIMIDO_MAXIMO:
                raise ArchivoInvalido('El archivo es demasiado grande para procesarse.')
    except zipfile.BadZipFile:
        raise ArchivoInvalido('El archivo no es un Excel (.xlsx) válido.')
    archivo.seek(0)

    try:
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ArchivoInvalido('No se pudo leer el archivo. Guárdalo de nuevo como .xlsx e inténtalo otra vez.')
    try:
        renglones = libro.worksheets[0].iter_rows(max_col=20, values_only=True)
        columnas = {}
        for indice, valor in enumerate(next(renglones, None) or ()):
            encabezado = _encabezado(valor)
            if not encabezado:
                continue
            campo = ENCABEZADOS_ACEPTADOS.get(encabezado)
            if campo is None or campo in columnas:
                raise ArchivoInvalido(
                    f'La columna "{_celda(valor)}" no es de la plantilla. '
                    'Usa solo Nombre(s), Apellidos, Grado y Grupo.'
                )
            columnas[campo] = indice
        for campo, titulo in COLUMNAS_EXCEL:
            if campo not in columnas:
                raise ArchivoInvalido(f'Falta la columna "{titulo}". Descarga la plantilla y no cambies los encabezados.')

        filas = []
        for numero, valores in enumerate(renglones, start=2):
            valores = valores or ()
            fila = {campo: _celda(valores[i]) if i < len(valores) else '' for campo, i in columnas.items()}
            if not any(fila.values()):
                continue
            if len(filas) == MAX_FILAS_LOTE:
                raise ArchivoInvalido(f'El archivo tiene más de {MAX_FILAS_LOTE} alumnos. Divídelo en varios archivos.')
            filas.append({'fila': numero, **fila})
    finally:
        libro.close()

    if not filas:
        raise ArchivoInvalido('El archivo no tiene alumnos. Escribe una fila por alumno debajo de los encabezados.')
    return filas


def previsualizar_excel(archivo, escuela):
    """Vista previa de la carga: cada fila con su estado (lista o errores), sin guardar nada."""
    filas = leer_excel_alumnos(archivo)
    limpias, errores = validar_filas(filas, escuela, ROL_ALUMNO)
    por_fila = {}
    for e in errores:
        por_fila.setdefault(e['fila'], []).append(e['mensaje'])
    listas = {f['fila']: f for f in limpias}

    resultado = []
    for fila in filas:
        numero = fila['fila']
        if numero in listas:
            limpio = listas[numero]
            resultado.append({
                'fila': numero, 'nombre': limpio['nombre'],
                'apellidos': ' '.join(p for p in (limpio['aPaterno'], limpio['aMaterno']) if p),
                'grado': limpio['grado'], 'grupo': limpio['grupo'], 'errores': [],
            })
        else:
            resultado.append({**fila, 'errores': por_fila.get(numero, ['Fila con datos no válidos'])})
    return {'filas': resultado, 'listas': len(listas), 'con_errores': len(filas) - len(listas)}


def plantilla_excel():
    """Plantilla .xlsx: solo encabezados (sin ejemplos, para no cargar alumnos ficticios)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Alumnos'
    hoja.append([titulo for _, titulo in COLUMNAS_EXCEL])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color='1B5E20')
        celda.fill = PatternFill('solid', fgColor='E3F2E6')
    for letra, ancho in zip('ABCD', (28, 28, 10, 10)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = 'A2'
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ── Lista ─────────────────────────────────────────────────────────────────────

def usuarios_gestionables(escuela, excluir=None):
    """Docentes y alumnos de la escuela (nunca de otra), sin el propio Enlace."""
    qs = Usuario.objects.filter(escuela=escuela, tipousuario_id__in=ROLES_GESTIONABLES)
    if excluir is not None:
        qs = qs.exclude(pk=excluir.pk)
    return qs


def buscar_usuarios(escuela, q='', rol='todos', estado='todos', grado=None, grupo=None, excluir=None):
    """Lista filtrada: cada palabra de `q` debe aparecer en nombre, apellidos, correo o usuario."""
    qs = usuarios_gestionables(escuela, excluir).select_related('tipousuario')
    for palabra in (q or '').split()[:10]:
        qs = qs.filter(
            Q(nombre__icontains=palabra) | Q(aPaterno__icontains=palabra) | Q(aMaterno__icontains=palabra)
            | Q(email__icontains=palabra) | Q(usuario__icontains=palabra)
        )
    if rol in FILTROS_ROL:
        qs = qs.filter(tipousuario_id=FILTROS_ROL[rol])
    if estado in FILTROS_ESTADO:
        qs = qs.filter(is_active=FILTROS_ESTADO[estado])
    if grado:
        qs = qs.filter(grado=grado)
    if grupo:
        qs = qs.filter(grupo=grupo)
    # Alumnos primero (por grado y grupo), después docentes
    return qs.order_by('-tipousuario_id', 'grado', 'grupo', 'aPaterno', 'aMaterno', 'nombre', 'pk')


def _celda_segura(valor):
    """Evita que Excel interprete el contenido como fórmula."""
    texto = '' if valor is None else str(valor)
    return "'" + texto if texto[:1] in ('=', '+', '-', '@', '\t', '\r') else texto


def exportar_csv(usuarios):
    salida = io.StringIO()
    salida.write('﻿')
    escritor = csv.writer(salida)
    escritor.writerow(['Nombre(s)', 'Apellido paterno', 'Apellido materno', 'Tipo', 'Grado', 'Grupo',
                       'Usuario', 'Correo', 'Teléfono', 'Alta', 'Estado'])
    for u in usuarios:
        alta = timezone.localtime(u.fecha_alta).strftime('%d/%m/%Y') if u.fecha_alta else ''
        escritor.writerow([_celda_segura(v) for v in (
            u.nombre, u.aPaterno, u.aMaterno, u.tipousuario.nombre_tipo if u.tipousuario else '',
            u.grado, u.grupo, u.usuario, u.email, u.telefono, alta, 'Activo' if u.is_active else 'Inactivo',
        )])
    return salida.getvalue()


def contexto_tutorial():
    """Datos variables del tutorial de la plantilla, tomados de las constantes del servidor."""
    return {
        'max_filas_lote': MAX_FILAS_LOTE,
        'excel_tamano_maximo': EXCEL_TAMANO_MAXIMO,
        'excel_tamano_maximo_mb': EXCEL_TAMANO_MAXIMO // (1024 * 1024),
        'columnas_excel': [titulo for _, titulo in COLUMNAS_EXCEL],
    }
